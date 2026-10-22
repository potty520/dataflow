# -*- coding: utf-8 -*-
"""生成 Word 版操作手册（docs/操作手册.docx），与 Markdown/PDF 内容一致。

依赖：python-docx（DSH 自带 Python 已包含）。
用法：<bundled-python> deploy/cmds/build_manual_docx.py
"""
import io
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = r"C:\Users\Administrator\Documents\1\datagov"
MD = os.path.join(ROOT, "docs", "操作手册.md")
OUT = os.path.join(ROOT, "docs", "操作手册.docx")
DOCS = os.path.join(ROOT, "docs")


def shade(paragraph, color="F2F4F6"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("第 ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    run2 = p.add_run(" 页 · DGP 数据治理平台 操作手册")
    run2.font.size = Pt(8)


def inline_runs(paragraph, text, base_size=10.5):
    """处理 **粗体** 与 `代码`。"""
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        else:
            r = paragraph.add_run(tok)
        r.font.size = Pt(base_size)


def build():
    md = io.open(MD, encoding="utf-8").read()
    lines = md.split("\n")
    doc = Document()

    # 页面设置
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    add_page_number_footer(sec)

    # 默认字体
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(10.5)

    i = 0
    first_h2 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            if buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run("\n".join(buf))
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                r.font.size = Pt(8)
                shade(p)
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text)
                r.bold = True
                r.font.size = Pt(24)
                r.font.color.rgb = RGBColor(0x14, 0x54, 0x8C)
                p.paragraph_format.space_before = Pt(120)
                p.paragraph_format.space_after = Pt(18)
            elif level == 2:
                if not first_h2:
                    doc.add_page_break()
                first_h2 = False
                h = doc.add_heading(text, level=1)
                for r in h.runs:
                    r.font.color.rgb = RGBColor(0x14, 0x54, 0x8C)
            elif level == 3:
                h = doc.add_heading(text, level=2)
                for r in h.runs:
                    r.font.color.rgb = RGBColor(0x1D, 0x5F, 0x92)
            else:
                doc.add_heading(text, level=3)
            i += 1
            continue

        # 分割线
        if stripped in ("---", "***"):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "D8DEE4")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            p = doc.add_paragraph()
            inline_runs(p, " ".join(buf))
            p.paragraph_format.left_indent = Cm(0.4)
            shade(p, "F2F7FB")
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for idx, h in enumerate(header):
                hdr[idx].text = ""
                p = hdr[idx].paragraphs[0]
                inline_runs(p, h, 9)
                for r in p.runs:
                    r.bold = True
            for row in rows:
                cells = table.add_row().cells
                for idx, c in enumerate(row[:len(header)]):
                    cells[idx].text = ""
                    inline_runs(cells[idx].paragraphs[0], c, 9)
            doc.add_paragraph()
            continue

        # 列表
        if re.match(r"^\s*[-*]\s+", line):
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                text = re.sub(r"^\s*[-*]\s+", "", lines[i])
                p = doc.add_paragraph(style="List Bullet")
                inline_runs(p, text)
                i += 1
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                text = re.sub(r"^\s*\d+\.\s+", "", lines[i])
                p = doc.add_paragraph(style="List Number")
                inline_runs(p, text)
                i += 1
            continue

        # 图片
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", stripped)
        if m:
            img = os.path.join(DOCS, m.group(2).replace("/", "\\"))
            if os.path.exists(img):
                try:
                    doc.add_picture(img, width=Cm(16.4))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception as e:
                    doc.add_paragraph("[图片加载失败: %s]" % m.group(2))
            else:
                doc.add_paragraph("[缺失图片: %s]" % m.group(2))
            i += 1
            continue

        # 空行
        if stripped == "":
            i += 1
            continue

        # 段落
        buf = []
        while i < len(lines) and lines[i].strip() != "" and not re.match(
                r"^(#{1,4}\s|```|\||>\s|\s*[-*]\s+|\s*\d+\.\s+|-)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        if buf:
            p = doc.add_paragraph()
            inline_runs(p, " ".join(buf))
            p.paragraph_format.first_line_indent = Pt(0)

    doc.save(OUT)
    print("docx ->", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    build()
